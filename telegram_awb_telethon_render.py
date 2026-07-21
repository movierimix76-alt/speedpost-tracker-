"""
Telegram AWB Matcher — Render Deployment Version
===================================================
यह वर्ज़न Render.com पर Background Worker के रूप में चलाने के लिए है।
सारी जानकारी (API_ID, API_HASH, SESSION_STRING, GROUP, EXCEL_FILE)
Render के Environment Variables से आती है — इस फाइल में कुछ भी हार्डकोड
करने की ज़रूरत नहीं।

Excel फाइल Render के Persistent Disk (/data) पर रखी जाती है ताकि
सर्वर रीस्टार्ट होने पर भी डेटा न मिटे।
"""

import os
import re
import asyncio
import logging
from openpyxl import load_workbook, Workbook
from telethon import TelegramClient, events
from telethon.sessions import StringSession

# ========================= Environment Variables से CONFIG =========================
API_ID = int(os.environ["API_ID"])
API_HASH = os.environ["API_HASH"]
SESSION_STRING = os.environ["SESSION_STRING"]
GROUP = os.environ["GROUP"]                              # जैसे "@mygroup" या numeric ID
EXCEL_FILE = os.environ.get("EXCEL_FILE", "/data/awb.xlsx")
SHEET_NAME = os.environ.get("SHEET_NAME") or None
FOUND_TEXT = os.environ.get("FOUND_TEXT", "मिल गया")
OLD_MESSAGES_LIMIT = int(os.environ.get("OLD_MESSAGES_LIMIT", "100000"))
# =====================================================================================

AWB_PATTERN = re.compile(r"\b[A-Z]{2}\d{9}IN\b")

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(message)s")
log = logging.getLogger(__name__)

client = TelegramClient(StringSession(SESSION_STRING), API_ID, API_HASH)


def ensure_excel_exists():
    """अगर persistent disk पर अभी तक Excel फाइल नहीं है, तो एक खाली फाइल बना देता है
    (header: AWB, Status)। अगर आपकी असली AWB लिस्ट वाली फाइल पहले से Render disk पर
    अपलोड है तो यह फंक्शन कुछ नहीं करेगा।"""
    if not os.path.exists(EXCEL_FILE):
        log.info("Excel फाइल नहीं मिली, नई खाली फाइल बनाई जा रही है...")
        wb = Workbook()
        ws = wb.active
        ws.append(["AWB", "Status"])
        wb.save(EXCEL_FILE)


def mark_found_in_excel(awb_set):
    if not awb_set:
        return []
    wb = load_workbook(EXCEL_FILE)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.active

    matched = []
    for row in ws.iter_rows(min_row=2):
        cell_a, cell_b = row[0], row[1]
        if cell_a.value is None:
            continue
        val = str(cell_a.value).strip()
        if val in awb_set and (cell_b.value or "").strip() != FOUND_TEXT:
            cell_b.value = FOUND_TEXT
            matched.append(val)

    if matched:
        wb.save(EXCEL_FILE)
    return matched


async def scan_old_messages():
    log.info(f"पुराने {OLD_MESSAGES_LIMIT} मैसेज पढ़े जा रहे हैं...")
    all_awbs = set()
    count = 0
    async for msg in client.iter_messages(GROUP, limit=OLD_MESSAGES_LIMIT):
        count += 1
        if msg.text:
            all_awbs.update(AWB_PATTERN.findall(msg.text))
        if count % 5000 == 0:
            log.info(f"...{count} मैसेज पढ़ लिए, अब तक {len(all_awbs)} अलग AWB मिले")

    log.info(f"कुल {count} मैसेज पढ़े, {len(all_awbs)} अलग-अलग AWB मिले। मिलान किया जा रहा है...")
    matched = mark_found_in_excel(all_awbs)
    log.info(f"पुराने मैसेज से {len(matched)} AWB मार्क हुए।")


@client.on(events.NewMessage())
async def on_new_message(event):
    chat = await event.get_chat()
    chat_ok = (
        str(getattr(chat, "username", "")) == GROUP.lstrip("@")
        or str(getattr(chat, "id", "")) == GROUP
        or f"-100{getattr(chat, 'id', '')}" == GROUP
    )
    if not chat_ok:
        return

    text = event.raw_text or ""
    found = set(AWB_PATTERN.findall(text))
    if not found:
        return
    matched = mark_found_in_excel(found)
    if matched:
        log.info(f"नया मैसेज — मार्क हुआ: {matched}")


async def main():
    ensure_excel_exists()
    await client.start()
    log.info("Telegram से कनेक्ट हो गया (Render पर चल रहा है)।")

    await scan_old_messages()

    log.info("अब नए मैसेज लाइव सुने जा रहे हैं...")
    await client.run_until_disconnected()


if __name__ == "__main__":
    asyncio.run(main())
