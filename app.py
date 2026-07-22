"""
Telegram AWB Matcher — MongoDB Persistent Version (v5)
=========================================================
बदलाव (v5):
- अब डेटा local Excel फाइल की जगह MongoDB में सेव होता है — इसलिए
  Render restart/redeploy होने पर भी डेटा नहीं खोता।
- Excel अभी भी /upload से अपलोड कर सकते हैं (उसमें से AWB नंबर
  पढ़कर MongoDB में डाले जाते हैं), और /download से हमेशा ताज़ा
  Excel फाइल MongoDB के डेटा से बनाकर डाउनलोड कर सकते हैं।
"""

import os
import re
import io
import asyncio
import logging
import threading
from datetime import datetime, timezone
from openpyxl import load_workbook, Workbook
from flask import Flask, send_file, request
from telethon import TelegramClient, events
from telethon.sessions import StringSession
from pymongo import MongoClient
from pymongo.errors import PyMongoError

# ========================= Environment Variables से CONFIG =========================
API_ID = int(os.environ["API_ID"])
API_HASH = os.environ["API_HASH"]
SESSION_STRING = os.environ["SESSION_STRING"]
GROUP = os.environ["GROUP"]
MONGODB_URI = os.environ["MONGODB_URI"]
DB_NAME = os.environ.get("DB_NAME", "awb_tracker")
COLLECTION_NAME = os.environ.get("COLLECTION_NAME", "awbs")
FOUND_TEXT = os.environ.get("FOUND_TEXT", "मिल गया")
OLD_MESSAGES_LIMIT = int(os.environ.get("OLD_MESSAGES_LIMIT", "100000"))
PORT = int(os.environ.get("PORT", "10000"))
# =====================================================================================

AWB_PATTERN = re.compile(r"\b[A-Za-z]{2}\d{9}[Ii][Nn]\b")

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(message)s")
log = logging.getLogger(__name__)

# ---------------------------- MongoDB ----------------------------
mongo_client = MongoClient(MONGODB_URI)
db = mongo_client[DB_NAME]
awbs_col = db[COLLECTION_NAME]
awbs_col.create_index("awb", unique=True)

app = Flask(__name__)

telegram_loop = None
telegram_target = None
telegram_ready = threading.Event()


def extract_awbs(text: str):
    return {m.upper() for m in AWB_PATTERN.findall(text or "")}


# ---------------------------- Web रूट्स ----------------------------
@app.route("/")
def health():
    total = awbs_col.count_documents({})
    found = awbs_col.count_documents({"status": FOUND_TEXT})
    return f"""
    <h2>AWB बॉट चल रहा है ✅</h2>
    <p>कुल AWB: {total} | मिल गया: {found}</p>
    <p><a href="/upload">यहां से अपनी असली AWB Excel फाइल अपलोड करें</a></p>
    <p><a href="/download">लेटेस्ट Excel फाइल डाउनलोड करें</a></p>
    """


@app.route("/download")
def download():
    wb = Workbook()
    ws = wb.active
    ws.append(["AWB", "Status"])
    for doc in awbs_col.find({}, {"_id": 0, "awb": 1, "status": 1}).sort("awb", 1):
        ws.append([doc["awb"], doc.get("status") or ""])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="awb.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/upload", methods=["GET", "POST"])
def upload():
    if request.method == "POST":
        f = request.files.get("file")
        if not f or not f.filename.endswith(".xlsx"):
            return "कृपया .xlsx फाइल चुनें। <a href='/upload'>वापस जाएं</a>"

        wb = load_workbook(f, data_only=True)
        ws = wb.active

        # नई शीट अपलोड होते ही पुराना पूरा डेटा साफ कर दिया जाता है,
        # ताकि पुरानी और नई शीट के AWB आपस में मिक्स न हों — हर बार
        # सिर्फ ताज़ा अपलोड की गई शीट का डेटा रहेगा।
        awbs_col.delete_many({})
        log.info("पुराना डेटा साफ किया गया, नई शीट डाली जा रही है...")

        added = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            val = row[0]
            if val is None or str(val).strip() == "":
                continue
            awb = str(val).strip().upper()
            awbs_col.update_one(
                {"awb": awb},
                {"$setOnInsert": {"awb": awb, "status": None,
                                   "created_at": datetime.now(timezone.utc)}},
                upsert=True,
            )
            added += 1
        log.info(f"नई Excel अपलोड — {added} AWB MongoDB में डाले गए। Rescan शुरू हो रहा है...")

        if telegram_ready.is_set() and telegram_loop and telegram_target:
            asyncio.run_coroutine_threadsafe(
                scan_old_messages(telegram_target), telegram_loop
            )
            return (
                f"फाइल अपलोड हो गई ✅ ({added} AWB डाले गए) — पुराने मैसेज दोबारा स्कैन हो रहे हैं। "
                "<a href='/'>होम पर जाएं</a>"
            )
        return f"फाइल अपलोड हो गई ({added} AWB), लेकिन बॉट अभी तैयार नहीं है। <a href='/'>होम पर जाएं</a>"

    return """
    <h3>अपनी AWB Excel फाइल अपलोड करें</h3>
    <p>ध्यान रखें: column A में AWB नंबर होने चाहिए (header row 1 में)</p>
    <form method="post" enctype="multipart/form-data">
      <input type="file" name="file" accept=".xlsx">
      <button type="submit">अपलोड करें</button>
    </form>
    """


# ---------------------------- MongoDB हेल्पर ----------------------------
def mark_found(awb_set):
    """awb_set में मौजूद AWB को MongoDB में ढूंढकर status='मिल गया' सेट करता है।
    सिर्फ वही AWB मार्क होंगे जो पहले से collection में मौजूद हैं (यानी आपकी
    अपलोड की शीट में थे)। लौटाता है कि कितने नए मार्क हुए।"""
    if not awb_set:
        return 0
    result = awbs_col.update_many(
        {"awb": {"$in": list(awb_set)}, "status": {"$ne": FOUND_TEXT}},
        {"$set": {"status": FOUND_TEXT, "found_at": datetime.now(timezone.utc)}},
    )
    return result.modified_count


# ---------------------------- Telegram लॉजिक ----------------------------
client = TelegramClient(StringSession(SESSION_STRING), API_ID, API_HASH)


async def get_target_entity():
    try:
        return await client.get_entity(GROUP)
    except Exception:
        log.info("सीधे entity नहीं मिली, dialogs sync किए जा रहे हैं...")
        async for dialog in client.iter_dialogs():
            entity = dialog.entity
            if (
                str(getattr(entity, "username", "")) == str(GROUP).lstrip("@")
                or str(entity.id) == str(GROUP)
                or f"-100{entity.id}" == str(GROUP)
            ):
                return entity
        raise ValueError(
            f"चैनल/ग्रुप '{GROUP}' नहीं मिला — पक्का करें कि आपका अकाउंट उसमें join है।"
        )


async def scan_old_messages(target):
    log.info(f"पुराने {OLD_MESSAGES_LIMIT} मैसेज पढ़े जा रहे हैं...")
    all_awbs = set()
    count = 0
    async for msg in client.iter_messages(target, limit=OLD_MESSAGES_LIMIT):
        count += 1
        if msg.text:
            all_awbs.update(extract_awbs(msg.text))
        if count % 5000 == 0:
            log.info(f"...{count} मैसेज पढ़ लिए, अब तक {len(all_awbs)} अलग AWB मिले")

    log.info(f"कुल {count} मैसेज पढ़े, {len(all_awbs)} अलग-अलग AWB मिले। MongoDB में मिलान किया जा रहा है...")
    matched = mark_found(all_awbs)
    log.info(f"स्कैन पूरा — {matched} नए AWB मार्क हुए।")


def register_new_message_handler(target_id):
    @client.on(events.NewMessage(chats=target_id))
    async def on_new_message(event):
        found = extract_awbs(event.raw_text)
        if not found:
            return
        matched = mark_found(found)
        if matched:
            log.info(f"नया मैसेज — {matched} AWB मार्क हुए।")


async def telegram_main():
    global telegram_loop, telegram_target
    await client.start()
    log.info("Telegram से कनेक्ट हो गया।")

    target = await get_target_entity()
    telegram_target = target
    telegram_loop = asyncio.get_event_loop()
    log.info(f"टारगेट चैनल/ग्रुप मिल गया: {getattr(target, 'title', target.id)}")

    register_new_message_handler(target.id)
    telegram_ready.set()

    await scan_old_messages(target)

    log.info("अब नए मैसेज लाइव सुने जा रहे हैं...")
    await client.run_until_disconnected()


def run_telegram_in_thread():
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    try:
        loop.run_until_complete(telegram_main())
    except Exception as e:
        log.error(f"Telegram thread में गलती: {e}")


if __name__ == "__main__":
    try:
        mongo_client.admin.command("ping")
        log.info("MongoDB से कनेक्ट हो गया।")
    except PyMongoError as e:
        log.error(f"MongoDB कनेक्ट नहीं हो पाया: {e}")

    t = threading.Thread(target=run_telegram_in_thread, daemon=True)
    t.start()
    app.run(host="0.0.0.0", port=PORT)
